"""Calamine-first spreadsheet parsing with provenance and row accounting."""

from __future__ import annotations

import re
import time
import zipfile
from collections.abc import Iterable
from itertools import chain
from pathlib import Path, PurePosixPath
from typing import Any
from xml.etree.ElementTree import ParseError

from openpyxl.utils.exceptions import InvalidFileException
from python_calamine import CalamineError, CalamineWorkbook

from suseoro.ingestion.contracts import (
    DocumentRole,
    FieldWarning,
    ParsedField,
    ParsedRow,
    ParseResult,
    Provenance,
    RowStatus,
)
from suseoro.ingestion.detection import detect_file_type
from suseoro.ingestion.file_store import StoredFile
from suseoro.ingestion.mapping import canonical_field_for_header, infer_mapping
from suseoro.ingestion.safety import inspect_zip, safe_xml_from_bytes

PARSER_VERSION = "tabular-v1"
_SHA256_PATTERN = re.compile(r"[0-9a-f]{64}")
_FORMULA_MARKERS = ("=", "+", "-", "@")


class CalamineCompatibilityError(RuntimeError):
    pass


def _load_with_calamine(path: Path) -> list[tuple[str, Iterable[list[Any]]]]:
    try:
        workbook = CalamineWorkbook.from_path(path)
        try:
            # Formula references and provenance use absolute worksheet coordinates.
            return [
                (name, workbook.get_sheet_by_name(name).to_python(skip_empty_area=False))
                for name in workbook.sheet_names
            ]
        finally:
            workbook.close()
    except (CalamineError, OSError, ValueError) as error:
        raise CalamineCompatibilityError(str(error)) from error


def _load_with_openpyxl(
    path: Path, *, read_only: bool, data_only: bool
) -> list[tuple[str, Iterable[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        path,
        read_only=read_only,
        data_only=data_only,
        keep_links=False,
    )
    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _column_number(reference: str) -> int:
    letters = re.match(r"[A-Z]+", reference.upper())
    if not letters:
        return 0
    value = 0
    for character in letters.group(0):
        value = value * 26 + ord(character) - ord("A") + 1
    return value


def _relationship_id(element: Any) -> str | None:
    return next(
        (
            value
            for key, value in element.attrib.items()
            if key.endswith("}id") or key == "id"
        ),
        None,
    )


def _worksheet_targets(archive: zipfile.ZipFile) -> list[str]:
    workbook = safe_xml_from_bytes(archive.read("xl/workbook.xml"))
    relationships = safe_xml_from_bytes(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        relationship.attrib.get("Id"): relationship.attrib.get("Target")
        for relationship in relationships
        if relationship.attrib.get("Type", "").endswith("/worksheet")
    }
    ordered: list[str] = []
    for sheet in workbook.iter():
        if not (sheet.tag.endswith("}sheet") or sheet.tag == "sheet"):
            continue
        target = targets.get(_relationship_id(sheet))
        if not target:
            continue
        if target.startswith("/"):
            normalized = PurePosixPath(target.lstrip("/"))
        else:
            normalized = PurePosixPath("xl") / PurePosixPath(target)
        if ".." in normalized.parts:
            raise ValueError("unsafe worksheet relationship target")
        ordered.append(normalized.as_posix())
    return ordered


def _xlsx_formula_cells(path: Path) -> dict[tuple[int, int, int], str]:
    found: dict[tuple[int, int, int], str] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            for sheet_number, name in enumerate(_worksheet_targets(archive), start=1):
                if name not in names:
                    raise KeyError(name)
                root = safe_xml_from_bytes(archive.read(name))
                for cell in root.iter():
                    if not cell.tag.endswith("}c") and cell.tag != "c":
                        continue
                    formula = next(
                        (
                            child
                            for child in cell
                            if child.tag.endswith("}f") or child.tag == "f"
                        ),
                        None,
                    )
                    if formula is not None:
                        reference = cell.attrib.get("r", "")
                        row_match = re.search(r"(\d+)$", reference)
                        if row_match:
                            expression = formula.text or ""
                            found[
                                (
                                    sheet_number,
                                    int(row_match.group(1)),
                                    _column_number(reference),
                                )
                            ] = "=" + expression
    except (OSError, zipfile.BadZipFile, KeyError, ValueError):
        return {}
    return found


def _nonblank(row: list[Any]) -> bool:
    return any(value not in (None, "") for value in row)


def _trim(row: list[Any]) -> list[Any]:
    values = list(row)
    while values and values[-1] in (None, ""):
        values.pop()
    return values


def _unique_headers(row: list[Any]) -> list[str]:
    counts: dict[str, int] = {}
    headers: list[str] = []
    for column, value in enumerate(row, start=1):
        base = str(value).strip() if value not in (None, "") else f"열{column}"
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}__{counts[base]}")
    return headers


def _find_header(rows: list[list[Any]]) -> int | None:
    best_index: int | None = None
    best_score = 0
    for index, row in enumerate(rows[:25]):
        headers = [str(value) if value is not None else "" for value in row]
        inferred = infer_mapping(headers, [])
        score = sum(field is not None for field in inferred.mapping.values())
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def _error_row(
    *,
    sheet: str | None,
    source_row: int,
    code: str,
    message: str,
    source_file_sha256: str,
) -> ParsedRow:
    return ParsedRow(
        status=RowStatus.ROW_ERROR,
        provenance=Provenance(
            sheet=sheet,
            source_row=source_row,
            source_file_sha256=source_file_sha256,
        ),
        raw_values={},
        fields={},
        error_code=code,
        error_message=message,
    )


def _parse_sheets(
    sheets: list[tuple[str, Iterable[list[Any]]]],
    *,
    role: DocumentRole,
    detected_format: str,
    parser_backend: str,
    encoding: str | None = None,
    source_kind: str = "FILE",
    formula_cells: dict[tuple[int, int, int], str] | None = None,
    started_at: float | None = None,
    source_file_sha256: str,
) -> ParseResult:
    rows_out: list[ParsedRow] = []
    header_rows: dict[str, int] = {}
    formulas = formula_cells or {}
    for sheet_number, (sheet_name, source_rows) in enumerate(sheets, start=1):
        iterator = iter(source_rows)
        prefix: list[list[Any]] = []
        for _ in range(25):
            try:
                prefix.append(_trim(list(next(iterator))))
            except StopIteration:
                break
        header_index = _find_header(prefix)
        if header_index is None:
            for index, source_row in enumerate(chain(prefix, iterator), start=1):
                row = _trim(list(source_row))
                if _nonblank(row):
                    rows_out.append(
                        _error_row(
                            sheet=sheet_name,
                            source_row=index,
                            code="HEADER_NOT_FOUND",
                            message="Could not identify a semantic header row",
                            source_file_sha256=source_file_sha256,
                        )
                    )
            continue
        header_rows[sheet_name] = header_index + 1
        header_source = prefix[header_index]
        headers = _unique_headers(header_source)
        semantics: list[str | None] = []
        claimed: set[str] = set()
        for header in headers:
            semantic = canonical_field_for_header(header.split("__", 1)[0])
            if semantic in claimed:
                semantic = None
            if semantic:
                claimed.add(semantic)
            semantics.append(semantic)
        required = {field for field in ("isbn", "title") if field in claimed}
        for source_index, source_row in enumerate(
            chain(prefix[header_index + 1 :], iterator), start=header_index + 2
        ):
            row = _trim(list(source_row))
            if not _nonblank(row):
                continue
            if len(row) > len(headers):
                extended_headers = _unique_headers(
                    header_source + [None] * (len(row) - len(header_source))
                )
                headers.extend(extended_headers[len(headers) :])
                semantics.extend([None] * (len(headers) - len(semantics)))
            padded = row + [None] * max(0, len(headers) - len(row))
            raw_values = {header: padded[index] for index, header in enumerate(headers)}
            explicit_formula_columns: set[int] = set()
            column_warnings: dict[int, tuple[FieldWarning, ...]] = {}
            row_warnings: list[FieldWarning] = []
            for column, header in enumerate(headers, start=1):
                formula_source = formulas.get((sheet_number, source_index, column))
                if formula_source is not None:
                    raw_values[header] = formula_source
                    explicit_formula_columns.add(column)
                raw_value = raw_values[header]
                if isinstance(raw_value, str) and raw_value.startswith(
                    _FORMULA_MARKERS
                ):
                    warning = FieldWarning(
                        "FORMULA_LIKE_INPUT",
                        f"Formula-like source text preserved without execution in column {header}",
                    )
                    column_warnings[column] = (warning,)
                    row_warnings.append(warning)
            fields: dict[str, ParsedField] = {}
            source_columns: dict[str, int] = {}
            for column, (header, semantic) in enumerate(
                zip(headers, semantics), start=1
            ):
                if not semantic:
                    continue
                raw_value = raw_values[header]
                warnings = column_warnings.get(column, ())
                value = raw_value
                if column in explicit_formula_columns:
                    value = None
                if (
                    semantic in {"isbn", "registration_number", "call_number"}
                    and value is not None
                ):
                    if isinstance(value, float) and value.is_integer():
                        value = str(int(value))
                    else:
                        value = str(value)
                fields[semantic] = ParsedField(
                    value=value, raw_value=raw_value, warnings=warnings
                )
                source_columns[semantic] = column
            missing = [
                field
                for field in sorted(required)
                if fields.get(field) is None or fields[field].value in (None, "")
            ]
            decode_error = any(
                isinstance(value, str) and "\ufffd" in value
                for value in raw_values.values()
            )
            status = (
                RowStatus.ROW_ERROR if missing or decode_error else RowStatus.SUCCESS
            )
            error_code = (
                "DECODE_ERROR"
                if decode_error
                else "MISSING_REQUIRED_FIELD"
                if missing
                else None
            )
            error_message = (
                "Source row contains invalid bytes for the detected encoding"
                if decode_error
                else ("Missing: " + ", ".join(missing))
                if missing
                else None
            )
            rows_out.append(
                ParsedRow(
                    status=status,
                    provenance=Provenance(
                        sheet=sheet_name,
                        source_row=source_index,
                        source_columns=source_columns,
                        source_file_sha256=source_file_sha256,
                    ),
                    raw_values=raw_values,
                    fields=fields,
                    warnings=tuple(row_warnings),
                    error_code=error_code,
                    error_message=error_message,
                )
            )
    elapsed = time.perf_counter() - started_at if started_at is not None else 0.000001
    return ParseResult(
        role=role,
        detected_format=detected_format,
        parser_version=PARSER_VERSION,
        parser_backend=parser_backend,
        rows=rows_out,
        header_rows=header_rows,
        encoding=encoding,
        source_kind=source_kind,
        elapsed_seconds=max(elapsed, 0.000001),
    )


def _source_path_and_sha256(
    source: Path | StoredFile, sha256: str | None
) -> tuple[Path, str]:
    if isinstance(source, StoredFile):
        path = source.path
        digest = source.sha256
        if sha256 is not None and sha256 != digest:
            raise ValueError("sha256 does not match StoredFile")
    else:
        path = Path(source)
        digest = sha256 or ""
    if _SHA256_PATTERN.fullmatch(digest) is None:
        raise ValueError("sha256 must be 64 lowercase hexadecimal characters")
    return path, digest


def parse_tabular(
    source: Path | StoredFile,
    *,
    role: DocumentRole,
    sha256: str | None = None,
) -> ParseResult:
    started = time.perf_counter()
    path, source_file_sha256 = _source_path_and_sha256(source, sha256)
    detection = detect_file_type(path)
    if detection.format in {"CSV", "TSV", "TXT"}:
        from suseoro.ingestion.parsers.text import parse_delimited_file

        result = parse_delimited_file(
            path,
            role=role,
            source_file_sha256=source_file_sha256,
            detection=detection,
        )
        return ParseResult(
            **{**result.__dict__, "elapsed_seconds": time.perf_counter() - started}
        )
    formula_cells: dict[tuple[int, int, int], str] = {}
    if detection.format in {"XLSX", "XLSB", "ODS"}:
        try:
            inspect_zip(path)
        except ValueError as error:
            return ParseResult(
                role=role,
                detected_format=detection.format,
                parser_version=PARSER_VERSION,
                parser_backend="none",
                rows=[
                    _error_row(
                        sheet=None,
                        source_row=0,
                        code="WORKBOOK_ERROR",
                        message=str(error),
                        source_file_sha256=source_file_sha256,
                    )
                ],
                elapsed_seconds=time.perf_counter() - started,
            )
    if detection.format == "XLSX":
        formula_cells = _xlsx_formula_cells(path)
    try:
        sheets = _load_with_calamine(path)
        backend = "calamine"
    except CalamineCompatibilityError as calamine_error:
        if detection.format != "XLSX":
            return ParseResult(
                role=role,
                detected_format=detection.format,
                parser_version=PARSER_VERSION,
                parser_backend="calamine",
                rows=[
                    _error_row(
                        sheet=None,
                        source_row=0,
                        code="WORKBOOK_ERROR",
                        message=str(calamine_error),
                        source_file_sha256=source_file_sha256,
                    )
                ],
                elapsed_seconds=time.perf_counter() - started,
            )
        try:
            sheets = _load_with_openpyxl(path, read_only=True, data_only=True)
            backend = "openpyxl-read-only-data-only"
        except (
            InvalidFileException,
            KeyError,
            OSError,
            ParseError,
            ValueError,
            zipfile.BadZipFile,
        ) as fallback_error:
            return ParseResult(
                role=role,
                detected_format=detection.format,
                parser_version=PARSER_VERSION,
                parser_backend="openpyxl-read-only-data-only",
                rows=[
                    _error_row(
                        sheet=None,
                        source_row=0,
                        code="WORKBOOK_ERROR",
                        message=str(fallback_error),
                        source_file_sha256=source_file_sha256,
                    )
                ],
                elapsed_seconds=time.perf_counter() - started,
            )
    return _parse_sheets(
        sheets,
        role=role,
        detected_format=detection.format,
        parser_backend=backend,
        formula_cells=formula_cells,
        started_at=started,
        source_file_sha256=source_file_sha256,
    )
