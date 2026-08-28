"""Inspect and copy a v1 workspace without modifying its source tree."""

from __future__ import annotations

import hashlib
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

_WORKBOOK_SUFFIXES = {".xlsx", ".xls", ".xlsb", ".ods"}


@dataclass(frozen=True)
class V1SchoolReport:
    name: str
    source_path: Path
    catalog_candidate: Path | None
    catalog_candidate_row_count: int | None
    requires_catalog_confirmation: bool
    legacy_workspaces: tuple[Path, ...]
    catalog_caches: tuple[Path, ...]


@dataclass
class V1MigrationReport:
    source_root: Path
    schools: list[V1SchoolReport] = field(default_factory=list)
    school_count: int = 0
    legacy_workspace_count: int = 0
    catalog_cache_count: int = 0
    catalog_candidate_count: int = 0
    missing_count: int = 0
    duplicate_count: int = 0
    read_error_count: int = 0
    copied_count: int = 0
    skipped_existing_count: int = 0
    source_delete_recommended: bool = False
    source_retention_message: str = (
        "두 번의 실제 수서 회차를 검증하기 전에는 v1 원본을 삭제하지 마세요."
    )

    def to_dict(self) -> dict[str, object]:
        return {
            "source_root": str(self.source_root),
            "school_count": self.school_count,
            "legacy_workspace_count": self.legacy_workspace_count,
            "catalog_cache_count": self.catalog_cache_count,
            "catalog_candidate_count": self.catalog_candidate_count,
            "missing_count": self.missing_count,
            "duplicate_count": self.duplicate_count,
            "read_error_count": self.read_error_count,
            "copied_count": self.copied_count,
            "skipped_existing_count": self.skipped_existing_count,
            "source_delete_recommended": self.source_delete_recommended,
            "source_retention_message": self.source_retention_message,
            "schools": [
                {
                    "name": school.name,
                    "catalog_candidate": str(school.catalog_candidate)
                    if school.catalog_candidate
                    else None,
                    "catalog_candidate_row_count": school.catalog_candidate_row_count,
                    "requires_catalog_confirmation": school.requires_catalog_confirmation,
                    "legacy_workspaces": [
                        str(path) for path in school.legacy_workspaces
                    ],
                    "catalog_caches": [str(path) for path in school.catalog_caches],
                }
                for school in self.schools
            ],
        }


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _workbook_rows(path: Path) -> int:
    if path.suffix.casefold() != ".xlsx":
        if path.stat().st_size == 0:
            raise ValueError("empty workbook")
        return 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        populated = sum(
            1
            for row in worksheet.iter_rows(values_only=True)
            if any(value is not None for value in row)
        )
        return max(0, populated - 1)
    finally:
        workbook.close()


def _is_legacy(path: Path, school: Path) -> bool:
    relative = path.relative_to(school)
    if len(relative.parts) > 1 and any("차" in part for part in relative.parts[:-1]):
        return True
    stem = path.stem.casefold()
    return "검수" in stem or "최종" in stem


def inspect_v1(source_root: Path) -> V1MigrationReport:
    root = Path(source_root).resolve(strict=True)
    if not root.is_dir():
        raise ValueError("v1 원본은 폴더여야 합니다.")
    report = V1MigrationReport(source_root=root)
    seen: set[str] = set()
    for school_path in sorted(path for path in root.iterdir() if path.is_dir()):
        workbooks = sorted(
            path
            for path in school_path.rglob("*")
            if path.is_file() and path.suffix.casefold() in _WORKBOOK_SUFFIXES
        )
        caches = tuple(
            path for path in workbooks if path.name.casefold() == "catalog_master.xlsx"
        )
        legacy = tuple(
            path
            for path in workbooks
            if path not in caches and _is_legacy(path, school_path)
        )
        dls_files = [
            path
            for path in workbooks
            if path not in caches
            and path not in legacy
            and "dls" in path.stem.casefold()
        ]
        valid_dls: list[tuple[Path, int]] = []
        for path in dls_files:
            try:
                valid_dls.append((path, _workbook_rows(path)))
            except (OSError, ValueError, KeyError, BadZipFile):
                report.read_error_count += 1
        valid_dls.sort(
            key=lambda item: (re.findall(r"\d+", item[0].stem), item[0].name),
            reverse=True,
        )
        selected = valid_dls[0] if valid_dls else None
        categorized = list(legacy) + list(caches) + ([selected[0]] if selected else [])
        for path in categorized:
            digest = _digest(path)
            if digest in seen:
                report.duplicate_count += 1
            seen.add(digest)
        school = V1SchoolReport(
            name=school_path.name,
            source_path=school_path,
            catalog_candidate=selected[0] if selected else None,
            catalog_candidate_row_count=selected[1] if selected else None,
            requires_catalog_confirmation=selected is not None,
            legacy_workspaces=legacy,
            catalog_caches=caches,
        )
        report.schools.append(school)
        report.legacy_workspace_count += len(legacy)
        report.catalog_cache_count += len(caches)
        report.catalog_candidate_count += int(selected is not None)
        report.missing_count += int(selected is None)
    report.school_count = len(report.schools)
    return report


def _copy(source: Path, destination: Path, report: V1MigrationReport) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists() and _digest(destination) == _digest(source):
        report.skipped_existing_count += 1
        return
    if destination.exists():
        destination = destination.with_name(
            f"{destination.stem}-{_digest(source)[:8]}{destination.suffix}"
        )
        if destination.exists() and _digest(destination) == _digest(source):
            report.skipped_existing_count += 1
            return
    shutil.copy2(source, destination)
    report.copied_count += 1


def migrate_v1(source_root: Path, destination_root: Path) -> V1MigrationReport:
    report = inspect_v1(source_root)
    destination = Path(destination_root).resolve()
    try:
        destination.relative_to(report.source_root)
    except ValueError:
        pass
    else:
        raise ValueError("이전 대상은 v1 원본 폴더 밖에 두어야 합니다.")
    for school in report.schools:
        target = destination / "schools" / school.name
        if school.catalog_candidate:
            _copy(
                school.catalog_candidate,
                target / "catalog-candidates" / school.catalog_candidate.name,
                report,
            )
        for cache in school.catalog_caches:
            _copy(cache, target / "catalog-cache" / cache.name, report)
        for legacy in school.legacy_workspaces:
            relative = legacy.relative_to(school.source_path)
            _copy(legacy, target / "legacy-workspaces" / relative, report)
    return report
