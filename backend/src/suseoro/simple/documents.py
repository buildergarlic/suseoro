"""Conservative local document imports with editable, source-accounted previews."""

from __future__ import annotations

import csv
import hashlib
import os
import re
import subprocess
import unicodedata
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from suseoro.catalog.normalization import canonical_isbn13
from suseoro.ingestion.contracts import DocumentRole, RowStatus
from suseoro.ingestion.detection import detect_file_type
from suseoro.ingestion.parsers.docx import parse_docx
from suseoro.ingestion.parsers.hwp import parse_hwp
from suseoro.ingestion.parsers.hwpx import parse_hwpx
from suseoro.ingestion.parsers.tabular import parse_tabular
from suseoro.ingestion.safety import inspect_zip

MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_ROWS = 200_000
MAX_PAGES = 1000
_ISBN_PATTERN = re.compile(r"(?<!\d)(?:97[89][\s-]?(?:\d[\s-]?){9}\d|(?:\d[\s-]?){9}[\dXx])(?!\d)")
_TABLE_NODE = re.compile(r"(?P<table>.*table\[\d+\])/row\[(?P<row>\d+)\]/cell\[(?P<cell>\d+)\]$")


def _key(value: object) -> str:
    return re.sub(r"[^a-z0-9가-힣]", "", unicodedata.normalize("NFKC", str(value or "")).casefold())


ALIASES = {
    "title": ("title", "제목", "도서명", "자료명", "서명", "책제목", "품명", "도서명서명"),
    "author": ("author", "authors", "저자", "지은이", "글쓴이", "작가", "저자명", "저자역자"),
    "publisher": ("publisher", "출판사", "발행처", "출판", "발행사"),
    "isbn": ("isbn", "isbn13", "isbn10", "국제표준도서번호", "도서번호"),
    "price": ("price", "listprice", "정가", "가격", "도서정가", "소비자가", "표시가격", "단가"),
    "quantity": ("quantity", "수량", "권수", "부수", "신청갯수", "신청개수", "주문수량"),
    "category": ("category", "분류", "주제", "KDC", "분류기호"),
    "requester": ("requester", "요청자", "신청자", "신청자명", "추천자", "신청자ID"),
    "audience": ("audience", "대상", "대상학년", "권장학년", "학년", "학교급"),
    "priority": ("priority", "우선순위", "중요도"),
    "source": ("source", "출처", "목록출처", "추천기관", "기관", "추천목록", "목록명"),
    "note": ("note", "메모", "비고", "추천사유", "신청사유"),
    "published_date": ("publisheddate", "publicationdate", "발행일", "발행년", "출판일", "발행연도"),
}
_FIELDS = {_key(alias): name for name, aliases in ALIASES.items() for alias in aliases}


def canonical_header(value: object) -> str | None:
    """Shared import/export Korean column-name recognition."""
    return _FIELDS.get(_key(value))


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value).strip()


def _safe_raw(values: dict[str, Any]) -> dict[str, Any]:
    return {
        str(key): value if value is None or isinstance(value, (str, int, bool)) else _text(value)
        for key, value in values.items()
    }


def _integer(value: Any, *, minimum: int) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    string = re.sub(r"(?:KRW|원|₩|,|\s)", "", _text(value), flags=re.IGNORECASE)
    if not re.fullmatch(r"\d+(?:\.0+)?", string):
        return None
    try:
        decimal = Decimal(string)
        return int(decimal) if minimum <= decimal <= 2_000_000_000 else None
    except (InvalidOperation, ValueError):
        return None


def _mapping(headers: list[str], provided: dict[str, str] | None) -> dict[str, str]:
    result: dict[str, str] = {}
    for header in headers:
        field = canonical_header(header.split("__", 1)[0])
        if field and field not in result:
            result[field] = header
    if provided is not None:
        # Public shape is canonical field -> original header. Accepting the
        # reverse shape also keeps imported older mapping preferences usable.
        result = {}
        for field, header in provided.items():
            if field in ALIASES and header in headers:
                result[field] = header
            elif header in ALIASES and field in headers:
                result[header] = field
    return result


def _book(
    values: dict[str, Any], *, filename: str, sheet: str, row: int,
    mapping: dict[str, str] | None = None, warnings: list[str] | None = None,
    page: int | None = None, raw_text: str = "", uncertain: bool = False,
) -> dict[str, Any]:
    raw = _safe_raw(values)
    resolved = _mapping(list(raw), mapping)
    fields = {field: raw.get(header) for field, header in resolved.items()}
    notices = list(warnings or [])
    title = _text(fields.get("title"))
    isbn_input = _text(fields.get("isbn"))
    isbn = canonical_isbn13(isbn_input) if isbn_input else None
    price = _integer(fields.get("price"), minimum=0)
    quantity_input = fields.get("quantity")
    quantity = 1 if quantity_input in (None, "") else _integer(quantity_input, minimum=1)
    if not title:
        notices.append("도서명을 확인해 주세요.")
    if not isbn:
        notices.append("ISBN이 없거나 체크숫자가 맞지 않습니다. 정확한 ISBN을 확인해 주세요.")
    if price is None:
        notices.append("정가를 확인해 주세요. 가격이 없으면 발주 금액에 포함되지 않습니다.")
    if quantity is None:
        notices.append("수량을 읽지 못해 1권으로 표시했습니다. 원본 수량을 확인해 주세요.")
        quantity = 1
    if any(isinstance(value, str) and value.lstrip().startswith(("=", "+", "@")) for value in raw.values()):
        notices.append("수식 형태의 원문은 실행하지 않았습니다. 값을 확인해 주세요.")
    if uncertain:
        notices.append("자동으로 서지를 확정할 수 없습니다. 원문과 비교해 주세요.")
    origin = f"{filename} · {sheet} · {row}행"
    if page is not None:
        origin = f"{filename} · {page}쪽 · {row}행"
    institution = _text(fields.get("source"))
    source = f"{institution} · {origin}" if institution else origin
    priority = _text(fields.get("priority")).lower()
    priority = {"상": "high", "높음": "high", "중": "normal", "보통": "normal", "하": "low", "낮음": "low"}.get(priority, priority)
    provenance: dict[str, Any] = {"filename": filename, "sheet": sheet, "row": row}
    if page is not None:
        provenance["page"] = page
    return {
        "title": title, "author": _text(fields.get("author")),
        "publisher": _text(fields.get("publisher")), "isbn": isbn or isbn_input,
        "price": price, "quantity": quantity, "selected": True,
        "category": _text(fields.get("category")), "requester": _text(fields.get("requester")),
        "audience": _text(fields.get("audience")), "priority": priority if priority in {"high", "normal", "low"} else "normal",
        "source": source, "note": _text(fields.get("note")),
        "published_date": _text(fields.get("published_date")), "link": "",
        "needs_review": bool(notices), "warnings": list(dict.fromkeys(notices)),
        "raw_values": raw, "raw_text": raw_text or " | ".join(_text(value) for value in raw.values()),
        "provenance": provenance,
    }


def _unique_headers(values: list[Any]) -> list[str]:
    result: list[str] = []
    for index, value in enumerate(values, 1):
        base = _text(value) or f"열 {index}"
        header = base
        suffix = 2
        while header in result:
            header = f"{base}__{suffix}"
            suffix += 1
        result.append(header)
    return result


def _matrix(
    values: list[list[Any]], *, filename: str, sheet: str,
    mapping: dict[str, str] | None = None, page: int | None = None,
) -> tuple[list[dict], list[str]]:
    nonempty = [(index, row) for index, row in enumerate(values, 1) if any(_text(cell) for cell in row)]
    if not nonempty:
        return [], []
    scores = [(sum(canonical_header(cell) is not None for cell in row), index, row) for index, row in nonempty[:25]]
    score, header_number, header_values = max(scores, key=lambda item: (item[0], -item[1]))
    if not score:
        header_number, header_values = nonempty[0]
        # A one-column ISBN list has no header. Keep every line for confirmation.
        if len(header_values) == 1:
            return _lines("\n".join(_text(row[0]) for _, row in nonempty), filename=filename, sheet=sheet, page=page), []
        if any(canonical_isbn13(value) or _integer(value, minimum=0) is not None for value in header_values):
            # A data row containing ISBNs/numbers is not silently consumed as
            # an unknown header. Expose generic columns for manual mapping.
            header_number = 0
            header_values = [f"열 {index}" for index in range(1, max(len(row) for _, row in nonempty) + 1)]
    headers = _unique_headers(header_values)
    rows = []
    for source_row, values_row in nonempty:
        if source_row == header_number:
            continue
        if source_row < header_number:
            rows.extend(_lines(" | ".join(_text(value) for value in values_row), filename=filename, sheet=sheet, page=page, start_row=source_row))
            continue
        extended = headers + [f"열 {index}" for index in range(len(headers) + 1, len(values_row) + 1)]
        raw = {header: values_row[index] if index < len(values_row) else None for index, header in enumerate(extended)}
        rows.append(_book(raw, filename=filename, sheet=sheet, row=source_row, mapping=mapping, page=page, uncertain=not bool(_mapping(extended, mapping))))
    return rows, headers


def _lines(text: str, *, filename: str, sheet: str, page: int | None = None, start_row: int = 1) -> list[dict]:
    result = []
    for index, line in enumerate(text.splitlines(), start_row):
        line = line.strip()
        if not line:
            continue
        matches = list(_ISBN_PATTERN.finditer(line))
        valid = [(match.group(), canonical_isbn13(match.group())) for match in matches]
        valid = [(raw, isbn) for raw, isbn in valid if isbn]
        if valid:
            for _, isbn in valid:
                result.append(_book({"isbn": isbn, "note": line}, filename=filename, sheet=sheet, row=index, page=page, raw_text=line, uncertain=True))
        else:
            result.append(_book({"note": line}, filename=filename, sheet=sheet, row=index, page=page, raw_text=line, uncertain=True))
    return result


def _text_matrix(path: Path) -> list[list[Any]]:
    contents = path.read_bytes()
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            text = contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = contents.decode("utf-8", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
        return list(csv.reader(text.splitlines(), dialect=dialect))
    except csv.Error:
        return [[line] for line in text.splitlines()]


def _tabular(path: Path, filename: str, mapping: dict[str, str] | None, digest: str) -> dict:
    parsed = parse_tabular(path, role=DocumentRole.PURCHASE_REQUEST, sha256=digest)
    rows: list[dict] = []
    headers: list[str] = []
    unknown_sheets = {row.provenance.sheet for row in parsed.rows if row.error_code == "HEADER_NOT_FOUND"}
    replacements: dict[str, tuple[list[dict], list[str]]] = {}
    if unknown_sheets:
        if parsed.detected_format in {"CSV", "TSV", "TXT"}:
            matrices = [("텍스트", _text_matrix(path))]
        elif parsed.detected_format == "XLSX":
            from openpyxl import load_workbook

            workbook = load_workbook(path, data_only=False, read_only=True, keep_links=False)
            try:
                matrices = [(sheet.title, [list(row) for row in sheet.values]) for sheet in workbook if sheet.title in unknown_sheets]
            finally:
                workbook.close()
        else:
            from python_calamine import CalamineWorkbook

            workbook = CalamineWorkbook.from_path(str(path))
            try:
                matrices = [(name, workbook.get_sheet_by_name(name).to_python()) for name in workbook.sheet_names if name in unknown_sheets]
            finally:
                workbook.close()
        for name, values in matrices:
            converted, sheet_headers = _matrix(values, filename=filename, sheet=name, mapping=mapping)
            replacements[name] = (converted, sheet_headers)
    replaced = set()
    for row in parsed.rows:
        sheet_name = row.provenance.sheet
        if sheet_name in replacements:
            if sheet_name not in replaced:
                converted, sheet_headers = replacements[sheet_name]
                rows.extend(converted)
                headers.extend(header for header in sheet_headers if header not in headers)
                replaced.add(sheet_name)
            continue
        raw = dict(row.raw_values)
        notices = [warning.message for warning in row.warnings]
        if row.error_message:
            notices.append(row.error_message)
        headers.extend(header for header in raw if header not in headers)
        rows.append(_book(raw, filename=filename, sheet=sheet_name or "표", row=row.provenance.source_row, mapping=mapping, warnings=notices, uncertain=row.status == RowStatus.ROW_ERROR))
    return {"rows": rows, "headers": headers, "mapping": _mapping(headers, mapping), "warnings": []}


def _document(path: Path, filename: str, kind: str, mapping: dict[str, str] | None, digest: str) -> dict:
    parsed = {"DOCX": parse_docx, "HWPX": parse_hwpx, "HWP": parse_hwp}[kind](path, role=DocumentRole.PURCHASE_REQUEST, sha256=digest)
    tables: dict[tuple[str, str], dict[int, dict[int, str]]] = defaultdict(lambda: defaultdict(dict))
    hwp_cells: dict[tuple[str, str], list[Any]] = defaultdict(list)
    other: list[Any] = []
    for row in parsed.rows:
        match = _TABLE_NODE.match(str(row.raw_values.get("node", "")))
        sheet = (row.provenance.sheet or kind).split("#", 1)[0]
        if row.status == RowStatus.ROW_ERROR:
            other.append(row)
        elif match:
            tables[(sheet, match.group("table"))][int(match.group("row"))][int(match.group("cell"))] = _text(row.fields["text"].value)
        elif kind == "HWP" and row.raw_values.get("cell") is not None:
            hwp_cells[(sheet, str(row.raw_values.get("table", 1)))].append(row)
        else:
            other.append(row)
    rows, headers, warnings = [], [], []
    for (sheet, table), entries in tables.items():
        width = max(column for cells in entries.values() for column in cells)
        matrix = [[cells.get(column, "") for column in range(1, width + 1)] for _, cells in sorted(entries.items())]
        converted, found_headers = _matrix(matrix, filename=filename, sheet=f"{sheet} {table}", mapping=mapping)
        rows.extend(converted)
        headers.extend(header for header in found_headers if header not in headers)
    for (sheet, table), cells in hwp_cells.items():
        cells.sort(key=lambda row: int(row.raw_values.get("cell") or 0))
        texts = [_text(row.fields.get("text").value) for row in cells]
        candidates = [(sum(canonical_header(value) is not None for value in texts[:width]), width) for width in range(2, min(30, len(texts) // 2) + 1) if len(texts) % width == 0]
        score, width = max(candidates, key=lambda item: (item[0], -item[1]), default=(0, 0))
        if score >= 2:
            converted, found_headers = _matrix([texts[offset:offset + width] for offset in range(0, len(texts), width)], filename=filename, sheet=f"{sheet} 표 {table}", mapping=mapping)
            # HWP v5 exposes sequential cells; retain the uncertainty from
            # inferring the number of columns instead of inventing exact layout.
            for book in converted:
                book["needs_review"] = True
                book["warnings"].append("HWP 표의 열 구조를 추정했습니다. 원본 표와 비교해 주세요.")
            rows.extend(converted)
            headers.extend(header for header in found_headers if header not in headers)
        else:
            rows.extend(_lines("\n".join(texts), filename=filename, sheet=f"{sheet} 표 {table}"))
    for row in other:
        text = _text(row.fields["text"].value) if "text" in row.fields else _text(row.raw_values.get("text"))
        sheet = row.provenance.sheet or kind
        if row.error_message:
            message = f"{sheet}: {row.error_message}"
            warnings.append(message)
            rows.append(_book({"note": text}, filename=filename, sheet=sheet, row=row.provenance.source_row, warnings=[message], uncertain=True))
        elif text:
            rows.extend(_lines(text, filename=filename, sheet=sheet, start_row=row.provenance.source_row))
    return {"rows": rows, "headers": headers, "mapping": _mapping(headers, mapping), "warnings": warnings}


def _ocr(path: Path, page_number: int) -> str:
    executable = os.environ.get("SUSEORO_TESSERACT", "").strip()
    if not executable or not Path(executable).is_file():
        raise RuntimeError("OCR을 사용할 수 없습니다. 한국어 OCR이 포함된 설치본을 사용하거나 텍스트 PDF/Excel로 다시 저장해 주세요.")
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(path))
    try:
        page = document[page_number - 1]
        try:
            bitmap = page.render(scale=2)
            try:
                from io import BytesIO

                output = BytesIO()
                bitmap.to_pil().save(output, format="PNG")
                image = output.getvalue()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        document.close()
    options: dict[str, Any] = {"creationflags": subprocess.CREATE_NO_WINDOW} if os.name == "nt" else {}
    try:
        completed = subprocess.run([executable, "stdin", "stdout", "-l", "kor+eng", "--psm", "6"], input=image, capture_output=True, timeout=45, check=True, shell=False, **options)
    except (OSError, subprocess.SubprocessError) as error:
        raise RuntimeError("OCR을 완료하지 못했습니다. 한국어·영어 언어 자료와 원본 페이지를 확인해 주세요.") from error
    return completed.stdout.decode("utf-8", errors="replace")


def _pdf(path: Path, filename: str, mapping: dict[str, str] | None) -> dict:
    import pdfplumber

    rows, warnings, headers = [], [], []
    with pdfplumber.open(path) as document:
        if len(document.pages) > MAX_PAGES:
            raise ValueError(f"PDF는 {MAX_PAGES:,}쪽 이하로 나누어 주세요.")
        for number, page in enumerate(document.pages, 1):
            page_start = len(rows)
            used_ocr = False
            try:
                tables = page.find_tables()
                for table_index, table in enumerate(tables, 1):
                    values = table.extract()
                    converted, found_headers = _matrix(values, filename=filename, sheet=f"표 {table_index}", mapping=mapping, page=number)
                    rows.extend(converted)
                    headers.extend(header for header in found_headers if header not in headers)
                boxes = [table.bbox for table in tables]
                def outside_tables(obj):
                    x = (obj.get("x0", 0) + obj.get("x1", 0)) / 2
                    y = (obj.get("top", 0) + obj.get("bottom", 0)) / 2
                    return not any(left <= x <= right and top <= y <= bottom for left, top, right, bottom in boxes)
                remaining = page.filter(outside_tables).extract_text() or ""
                if not tables and len(remaining.strip()) < 12:
                    remaining = _ocr(path, number)
                    used_ocr = True
                    if not remaining.strip():
                        raise RuntimeError("OCR에서 글자를 읽지 못했습니다. 원본 페이지를 확인해 주세요.")
                    warnings.append(f"{number}쪽: OCR로 읽었습니다. ISBN과 가격을 원본과 비교해 주세요.")
                if remaining.strip():
                    nonempty = [line for line in remaining.splitlines() if line.strip()]
                    table_values = None
                    for delimiter in ("\t", "|", ";", ","):
                        candidate = list(csv.reader(nonempty, delimiter=delimiter))
                        if any(sum(canonical_header(cell) is not None for cell in line) >= 2 for line in candidate[:25]):
                            table_values = candidate
                            break
                    if table_values:
                        converted, found_headers = _matrix(table_values, filename=filename, sheet="본문", mapping=mapping, page=number)
                        rows.extend(converted)
                        headers.extend(header for header in found_headers if header not in headers)
                    else:
                        rows.extend(_lines(remaining, filename=filename, sheet="본문", page=number))
                if used_ocr:
                    for book in rows[page_start:]:
                        book["needs_review"] = True
                        book["warnings"].append("OCR로 읽은 ISBN과 가격을 원본과 비교해 주세요.")
            except Exception as error:
                # One corrupt or unreadable page never erases other pages.
                message = f"{number}쪽: {error}"
                warnings.append(message)
                rows.append(_book({"note": message}, filename=filename, sheet="PDF", row=1, page=number, warnings=[message], uncertain=True))
    return {"rows": rows, "headers": headers, "mapping": _mapping(headers, mapping), "warnings": warnings}


def parse_upload(path: Path, filename: str, mapping: dict[str, str] | None = None) -> dict:
    """Return editable Book rows and canonical-field -> source-header mapping.

    Files are never modified. Failed rows/pages are retained as review records;
    file-level safety violations raise ValueError for a visible upload error.
    """
    path = Path(path)
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("파일은 100MB 이하로 나누어 주세요.")
    with path.open("rb") as source:
        digest = hashlib.file_digest(source, "sha256").hexdigest()
    detected = detect_file_type(path)
    kind = detected.format
    if kind in {"XLSX", "XLSB", "ODS", "HWPX", "DOCX"}:
        inspect_zip(path)
    if kind in {"XLSX", "XLS", "XLSB", "ODS", "CSV", "TSV", "TXT"}:
        result = _tabular(path, filename, mapping, digest)
    elif kind in {"DOCX", "HWPX", "HWP"}:
        result = _document(path, filename, kind, mapping, digest)
    elif kind == "PDF":
        result = _pdf(path, filename, mapping)
    else:
        raise ValueError("Excel, CSV, TXT, PDF, HWP, HWPX, DOCX 파일을 선택해 주세요.")
    if len(result["rows"]) > MAX_ROWS:
        raise ValueError(f"한 번에 {MAX_ROWS:,}행까지 가져올 수 있습니다. 파일을 나누어 주세요.")
    for row in result["rows"]:
        row["provenance"]["sha256"] = digest
    if not result["rows"]:
        result["warnings"].append("읽을 수 있는 도서 행이 없습니다. 파일 내용과 열 제목을 확인해 주세요.")
    result["warnings"] = list(dict.fromkeys(result["warnings"]))
    return result
