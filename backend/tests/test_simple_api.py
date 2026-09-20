import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from suseoro.simple.app import create_app
from suseoro.simple import VERSION


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path), base_url='http://127.0.0.1:3847') as client:
        bootstrap = client.get('/api/library/bootstrap').json()
        client.headers['X-Suseoro-Token'] = bootstrap['csrf_token']
        yield client


def list_id(client):
    return client.get('/api/library/bootstrap').json()['lists'][0]['id']


def test_first_launch_needs_no_account(client):
    data = client.get('/api/library/bootstrap').json()
    assert data['version'] == VERSION
    assert data['lists'][0]['budget'] == 15_000_000
    assert not data['settings']['nl_api_key_configured']


def test_cross_origin_and_missing_token_cannot_change_data(client):
    assert client.post('/api/library/lists', json={'name': '외부 변경'}, headers={'Origin': 'https://evil.example'}).status_code == 403
    assert client.post('/api/library/lists', json={'name': '토큰 없음'}, headers={'X-Suseoro-Token': ''}).status_code == 403
    assert client.get('/api/library/bootstrap', headers={'Host': 'evil.example'}).status_code == 403


def test_add_edit_select_delete_restore(client):
    base = f'/api/library/lists/{list_id(client)}'
    response = client.post(base + '/books', json={'title': '교사 추천도서', 'price': 17000})
    assert response.status_code == 200
    book = response.json()
    assert client.patch(base + '/books/' + book['id'], json={'quantity': 2}).status_code == 200
    assert client.get(base).json()['summary']['order_total'] == 34000
    client.delete(base + '/books/' + book['id'])
    assert client.get(base).json()['books'] == []
    client.post(base + '/books/' + book['id'] + '/restore')
    assert len(client.get(base).json()['books']) == 1


def test_key_never_echoed_in_settings_or_backup(client):
    response = client.patch('/api/library/settings', json={'nl_api_key': 'private-key', 'school_name': '함께초'})
    assert response.json()['nl_api_key_configured'] is True
    assert 'private-key' not in client.get('/api/library/bootstrap').text
    assert 'private-key' not in client.get('/api/library/backup').text


@pytest.mark.parametrize('size', [16, 18, 20, 22, 24])
def test_display_preferences_save_through_api_and_survive_a_new_app_port(tmp_path, size):
    with TestClient(create_app(tmp_path), base_url='http://127.0.0.1:3847') as first:
        initial = first.get('/api/library/bootstrap').json()
        assert initial['settings']['text_size'] == 16
        assert initial['settings']['row_density'] == 'comfortable'
        response = first.patch('/api/library/settings', headers={'X-Suseoro-Token': initial['csrf_token']}, json={'text_size': size, 'row_density': 'compact'})
        assert response.status_code == 200
        assert response.json()['text_size'] == size
    with TestClient(create_app(tmp_path), base_url='http://127.0.0.1:51234') as reopened:
        settings = reopened.get('/api/library/bootstrap').json()['settings']
        assert settings['text_size'] == size
        assert settings['row_density'] == 'compact'


@pytest.mark.parametrize('patch', [{'text_size': 17}, {'text_size': 26}, {'text_size': '20'}, {'text_size': True}, {'row_density': 'dense'}, {'row_density': []}])
def test_invalid_display_preferences_are_rejected_by_api(client, patch):
    before = client.get('/api/library/bootstrap').json()['settings']
    response = client.patch('/api/library/settings', json=patch)
    assert response.status_code == 400
    assert client.get('/api/library/bootstrap').json()['settings'] == before


def test_excel_import_then_export_real_workbook(client):
    workbook = Workbook()
    workbook.active.append(['도서명', '저자', 'ISBN', '정가'])
    workbook.active.append(['채식주의자', '한강', '9788936434267', 15000])
    stream = io.BytesIO()
    workbook.save(stream)
    preview = client.post('/api/library/imports/preview', files={'file': ('추천도서.xlsx', stream.getvalue())})
    assert preview.status_code == 200, preview.text
    data = preview.json()
    assert len(data['rows']) == 1
    base = f'/api/library/lists/{list_id(client)}'
    committed = client.post(base + '/imports', json={'import_id': data['import_id'], 'rows': data['rows'], 'kind': 'recommendations'})
    assert committed.status_code == 200, committed.text
    assert committed.json()['added'] == 1
    download = client.post(base + '/export', json={'format': 'xlsx'})
    assert download.status_code == 200, download.text
    book = load_workbook(io.BytesIO(download.content))
    assert '채식주의자' in str(list(book.active.values))


def test_unknown_price_blocks_final_export(client):
    base = f'/api/library/lists/{list_id(client)}'
    client.post(base + '/books', json={'title': '가격 미정'})
    response = client.post(base + '/export', json={'format': 'xlsx'})
    assert response.status_code == 400
    assert '가격' in response.json()['detail']


def test_backup_restore_round_trip(client):
    base = f'/api/library/lists/{list_id(client)}'
    client.post(base + '/books', json={'title': '보관할 책', 'price': 12000})
    backup = client.get('/api/library/backup').content
    client.post(base + '/books', json={'title': '나중 책', 'price': 11000})
    restored = client.post('/api/library/restore', files={'file': ('backup.json', backup)})
    assert restored.status_code == 200, restored.text
    assert len(client.get(base).json()['books']) == 1


def test_unsupported_upload_is_rejected(client):
    assert client.post('/api/library/imports/preview', files={'file': ('malicious.exe', b'MZ')}).status_code == 400
