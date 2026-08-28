from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook


def test_dls_isbn_txt_is_exact_crlf_utf8_without_bom_or_final_newline() -> None:
    from suseoro.exports.dls import dls_isbn_txt

    output = dls_isbn_txt(
        ["978-89-374-6401-0", "9788937464010", "invalid", "9788936434267"]
    )
    assert output == b"9788937464010\r\n9788936434267"
    assert not output.startswith(b"\xef\xbb\xbf")
    assert not output.endswith(b"\r\n")


def test_dls_title_xlsx_has_exact_sheet_columns_defaults_and_types() -> None:
    from suseoro.exports.dls import dls_title_xlsx

    output = dls_title_xlsx(
        [
            {
                "title": "파친코",
                "author": "이민진",
                "publisher": "문학사상",
                "isbn": "9788937464010",
                "quantity": 2,
            },
            {"title": "제목만"},
        ]
    )
    workbook = load_workbook(BytesIO(output), data_only=False)
    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook["Sheet1"]
    assert [cell.value for cell in sheet[1]] == [
        "자료명",
        "저자",
        "출판사",
        "자료유형",
        "신청갯수",
        "신청자ID",
        "ISBN",
    ]
    assert [cell.value for cell in sheet[2]] == [
        "파친코",
        "이민진",
        "문학사상",
        "단행본",
        2,
        "suseoro",
        "9788937464010",
    ]
    assert [cell.value for cell in sheet[3]] == [
        "제목만",
        "미상",
        "미상",
        "단행본",
        1,
        "suseoro",
        "0",
    ]
    assert sheet["E2"].data_type == "n"
    assert sheet["F2"].data_type == sheet["G2"].data_type == "s"


def test_every_user_controlled_workbook_cell_is_formula_safe() -> None:
    from suseoro.exports.dls import dls_title_xlsx
    from suseoro.exports.orders import order_xlsx
    from suseoro.exports.safe_cells import escape_spreadsheet_cell

    dangerous = [
        "=1+1",
        "+SUM(A1:A2)",
        "-2+3",
        "@cmd",
        "\t=cmd",
        "\r=cmd",
        "\n=cmd",
        "\x01cmd",
    ]
    for value in dangerous:
        escaped = escape_spreadsheet_cell(value)
        assert escaped.startswith("'")
    dls = load_workbook(
        BytesIO(
            dls_title_xlsx(
                [
                    {
                        "title": dangerous[0],
                        "author": dangerous[1],
                        "publisher": dangerous[4],
                    }
                ]
            )
        ),
        data_only=False,
    )["Sheet1"]
    order = load_workbook(
        BytesIO(
            order_xlsx(
                [
                    {
                        "isbn": "9788937464010",
                        "title": dangerous[2],
                        "author": dangerous[3],
                        "publisher": dangerous[5],
                        "quantity": 1,
                        "unit_price": 1000,
                    }
                ]
            )
        ),
        data_only=False,
    )["발주서"]
    for cell in (
        dls["A2"],
        dls["B2"],
        dls["C2"],
        order["B2"],
        order["C2"],
        order["D2"],
    ):
        assert cell.data_type == "s"
        assert cell.value.startswith("'")


def test_order_xlsx_default_and_template_column_order() -> None:
    from suseoro.exports.orders import DEFAULT_ORDER_COLUMNS, order_xlsx

    row = {
        "isbn": "9788937464010",
        "title": "책",
        "author": "저자",
        "publisher": "출판사",
        "quantity": 2,
        "unit_price": 9000,
    }
    default_sheet = load_workbook(BytesIO(order_xlsx([row])))["발주서"]
    assert [cell.value for cell in default_sheet[1]] == [
        label for _, label in DEFAULT_ORDER_COLUMNS
    ]
    assert default_sheet["G2"].value == 18_000
    template = (("title", "도서명"), ("isbn", "ISBN"), ("quantity", "수량"))
    template_sheet = load_workbook(BytesIO(order_xlsx([row], columns=template)))[
        "발주서"
    ]
    assert [cell.value for cell in template_sheet[1]] == ["도서명", "ISBN", "수량"]
    assert [cell.value for cell in template_sheet[2]] == ["책", "9788937464010", 2]


def test_user_controlled_template_headers_are_formula_safe() -> None:
    from suseoro.exports.orders import order_xlsx

    workbook = load_workbook(
        BytesIO(
            order_xlsx([{"title": "책"}], columns=(("title", '=HYPERLINK("bad")'),))
        ),
        data_only=False,
    )
    header = workbook["발주서"]["A1"]
    assert header.data_type == "s"
    assert header.value == '\'=HYPERLINK("bad")'
