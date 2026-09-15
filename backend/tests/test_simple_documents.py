from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import Workbook

from suseoro.simple.documents import parse_upload


def test_spreadsheet_keeps_invalid_and_missing_rows_with_provenance(tmp_path: Path):
    path = tmp_path / "기관 추천.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "중학교"
    sheet.append(["도서명", "지은이", "출판사", "ISBN", "정가", "권수", "추천기관"])
    sheet.append(["책 하나", "김작가", "출판", "9788937464010", "15,000원", 2, "어린이도서연구회"])
    sheet.append(["가격 없음", "저자", "출판", "9788936434267", None, 1, "기관"])
    sheet.append([None, "저자", "출판", "9788937464011", "=1+1", -1, "기관"])
    workbook.save(path)
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 3
    first, missing, invalid = result["rows"]
    assert first["price"] == 15000 and first["quantity"] == 2
    assert first["author"] == "김작가" and not first["needs_review"]
    assert first["provenance"]["row"] == 2
    assert first["provenance"]["sheet"] == "중학교"
    assert "어린이도서연구회" in first["source"]
    assert missing["price"] is None and missing["needs_review"]
    assert invalid["price"] is None and invalid["needs_review"]
    assert invalid["raw_values"]["정가"] == "=1+1"


def test_unknown_csv_headers_are_preserved_and_can_be_mapped(tmp_path: Path):
    path = tmp_path / "custom.csv"
    path.write_text("A,B,C\n책,9788937464010,12000\n다른 책,9788936434267,9000", encoding="utf-8")
    preview = parse_upload(path, path.name)
    assert preview["headers"] == ["A", "B", "C"]
    assert len(preview["rows"]) == 2
    assert preview["rows"][0]["raw_values"]["A"] == "책"
    mapped = parse_upload(path, path.name, mapping={"A": "title", "B": "isbn", "C": "price"})
    assert mapped["rows"][0]["title"] == "책"
    assert mapped["rows"][0]["price"] == 12000
    assert not mapped["rows"][0]["needs_review"]


def test_hwpx_explicit_table_is_reconstructed(tmp_path: Path):
    path = tmp_path / "추천.hwpx"
    values = [["도서명", "ISBN", "정가"], ["책", "9788937464010", "10000"], ["다른 책", "9788936434267", ""]]
    xml = '<root xmlns:hp="urn:hwp"><hp:tbl>' + "".join(
        "<hp:tr>" + "".join(f"<hp:tc><hp:p><hp:run><hp:t>{v}</hp:t></hp:run></hp:p></hp:tc>" for v in row) + "</hp:tr>"
        for row in values
    ) + "</hp:tbl></root>"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/hwp+zip")
        archive.writestr("Contents/section0.xml", xml)
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 2
    assert result["rows"][0]["title"] == "책"
    assert result["rows"][0]["price"] == 10000
    assert result["rows"][1]["needs_review"]


def test_pdf_without_ocr_keeps_page_as_review(tmp_path: Path, monkeypatch):
    from pypdf import PdfWriter

    monkeypatch.delenv("SUSEORO_TESSERACT", raising=False)
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["needs_review"]
    assert result["rows"][0]["provenance"]["page"] == 1
    assert any("OCR" in warning for warning in result["warnings"])


def test_plain_isbn_lines_are_not_given_invented_titles(tmp_path: Path):
    path = tmp_path / "메모.txt"
    path.write_text("추천도서\n9788937464010\n9788936434267", encoding="utf-8")
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 3
    isbn_rows = [row for row in result["rows"] if row["isbn"]]
    assert len(isbn_rows) == 2
    assert all(row["title"] == "" and row["needs_review"] for row in isbn_rows)


def test_pdf_positioned_table_is_parsed_without_ocr(tmp_path: Path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    path = tmp_path / "table.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=500, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    commands = ["0.5 w"]
    for y in (200, 225, 250):
        commands.append(f"40 {y} m 450 {y} l S")
    for x in (40, 180, 350, 450):
        commands.append(f"{x} 200 m {x} 250 l S")
    for x, y, text in [(45, 234, "title"), (185, 234, "ISBN"), (355, 234, "price"), (45, 209, "A book"), (185, 209, "9788937464010"), (355, 209, "12000")]:
        commands.append(f"BT /F1 10 Tf {x} {y} Td ({text}) Tj ET")
    stream = DecodedStreamObject()
    stream.set_data("\n".join(commands).encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["title"] == "A book"
    assert result["rows"][0]["price"] == 12000
    assert not result["rows"][0]["needs_review"]


def test_workbook_unknown_sheet_rows_preserve_values(tmp_path: Path):
    path = tmp_path / "mixed.xlsx"
    workbook = Workbook()
    workbook.active.append(["도서명", "ISBN", "정가"])
    workbook.active.append(["책", "9788937464010", 10000])
    sheet = workbook.create_sheet("unknown")
    sheet.append(["A", "B", "C"])
    sheet.append(["숨겨지면 안 되는 책", "9788936434267", 9000])
    workbook.save(path)
    result = parse_upload(path, path.name)
    unknown = [book for book in result["rows"] if book["provenance"]["sheet"] == "unknown"]
    assert len(unknown) == 1
    assert unknown[0]["raw_values"]["A"] == "숨겨지면 안 되는 책"


def test_hwp_cells_are_conservatively_reconstructed_and_errors_kept(tmp_path: Path):
    from test_hwp_parser import _minimal_hwp

    path = tmp_path / "추천.hwp"
    path.write_bytes(_minimal_hwp(unsupported=True, cell_values=["도서명", "ISBN", "정가", "책 하나", "9788937464010", "12000", "책 둘", "9788936434267", "9000", "책 셋", "9788937464010", "8000"]))
    result = parse_upload(path, path.name)
    books = [book for book in result["rows"] if book["isbn"]]
    assert len(books) == 3
    assert [book["title"] for book in books] == ["책 하나", "책 둘", "책 셋"]
    assert all(book["needs_review"] for book in books)
    assert result["warnings"]


def test_docx_table_cells_form_one_book_row(tmp_path: Path):
    path = tmp_path / "추천.docx"
    matrix = [["도서명", "ISBN", "정가"], ["책", "9788937464010", "12000"]]
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:tbl>' + "".join(
        "<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{value}</w:t></w:r></w:p></w:tc>" for value in row) + "</w:tr>" for row in matrix
    ) + "</w:tbl></w:body></w:document>"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["title"] == "책"
    assert result["rows"][0]["price"] == 12000


def test_ocr_result_requires_review_even_when_table_fields_are_complete(tmp_path: Path, monkeypatch):
    from pypdf import PdfWriter

    path = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(path)
    monkeypatch.setattr("suseoro.simple.documents._ocr", lambda *args: "title|ISBN|price\nA book|9788937464010|10000")
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 1
    assert result["rows"][0]["price"] == 10000
    assert result["rows"][0]["needs_review"]
    assert any("OCR" in item for item in result["rows"][0]["warnings"])


def test_headerless_csv_does_not_consume_first_book_as_header(tmp_path: Path):
    path = tmp_path / "no-header.csv"
    path.write_text("책 하나,9788937464010,10000\n책 둘,9788936434267,9000", encoding="utf-8")
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 2
    assert result["headers"] == ["열 1", "열 2", "열 3"]
    remapped = parse_upload(path, path.name, mapping={"title": "열 1", "isbn": "열 2", "price": "열 3"})
    assert remapped["rows"][0]["title"] == "책 하나"
    assert not remapped["rows"][0]["needs_review"]
