from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import pytest

from suseoro.simple.exporting import export_books, inspect_template


BOOKS = [
    {"title": "책 하나", "author": "김작가", "publisher": "출판", "isbn": "9788937464010", "price": 10005, "quantity": 2, "selected": True},
    {"title": "=1+1", "author": "이작가", "publisher": "출판", "isbn": "9788936434267", "price": 9000, "quantity": 1, "selected": True},
]
LIST = {"name": "2학기 구입", "year": 2026, "budget": 15000000, "discount_percent": 10}


def test_merged_recommendation_sources_export_together_without_changing_quantity():
    book = {**BOOKS[0], "isbn": "", "source": "수정한 출처", "sources": ["기관 A", "기관 B", "기관 A"]}
    content, _, _ = export_books([book], LIST, "학교")
    sheet = load_workbook(BytesIO(content)).active
    headers = {cell.value: cell.column for cell in sheet[6]}
    assert sheet.cell(7, headers["목록출처"]).value == "수정한 출처 · 기관 A · 기관 B"
    assert sheet.cell(7, headers["수량"]).value == 2
    assert sheet.cell(7, headers["ISBN"]).value is None


def test_long_merged_sources_are_preserved_in_separate_sheet():
    sources = [f"기관 {i:03} " + "가" * 100 for i in range(500)]
    content, _, _ = export_books([{**BOOKS[0], "source": sources[0], "sources": sources}], LIST, "학교")
    workbook = load_workbook(BytesIO(content))
    assert "추천출처" in workbook.sheetnames
    history = workbook["추천출처"]
    assert [history.cell(index + 2, 3).value for index in range(500)] == sources
    assert "추천출처" in workbook.active.cell(7, 12).value


def test_standard_workbook_uses_unit_half_up_discount_and_safe_text():
    content, filename, mime = export_books(BOOKS, LIST, "수서중학교")
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook.active
    cells = [cell for row in sheet for cell in row]
    assert any(cell.value == "수서중학교" for cell in cells)
    assert any(cell.value == 9005 for cell in cells)
    assert any(cell.value == 18010 for cell in cells)
    assert any(cell.value == "'=1+1" and cell.data_type == "s" for cell in cells)
    assert any(cell.value == "9788937464010" and cell.data_type == "s" for cell in cells)
    assert filename.endswith(".xlsx") and "spreadsheetml" in mime


def test_custom_school_template_preserves_heading_merge_style_and_footer(tmp_path: Path):
    path = tmp_path / "학교 양식.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "신청서"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "{{school_name}} 도서구입 신청서"
    sheet.append([])
    sheet.append(["번호", "도서명", "ISBN", "정가", "수량", "금액"])
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="FFEEDD")
    sheet["A5"] = "담당자 확인"
    sheet.merge_cells("A5:F5")
    workbook.save(path)
    info = inspect_template(path)
    assert "도서명" in info["columns"]
    content, _, _ = export_books(BOOKS, LIST, "수서중학교", template_path=path)
    sheet = load_workbook(BytesIO(content)).active
    assert sheet["A1"].value == "수서중학교 도서구입 신청서"
    assert sheet["B4"].value == "책 하나"
    assert sheet["B5"].value == "'=1+1"
    assert sheet["B5"].fill.fgColor.rgb == "00FFEEDD"
    assert sheet["A6"].value == "담당자 확인"
    assert "A6:F6" in {str(value) for value in sheet.merged_cells.ranges}
    assert path.exists() and load_workbook(path).active["A5"].value == "담당자 확인"


def test_placeholder_template_and_safe_csv_html(tmp_path: Path):
    path = tmp_path / "placeholder.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["{{title}}", "{{isbn}}", "{{quantity}}", "{{line_total}}"])
    workbook.save(path)
    content, _, _ = export_books(BOOKS, LIST, "학교", template_path=path)
    sheet = load_workbook(BytesIO(content)).active
    assert sheet["A1"].value == "책 하나" and sheet["A2"].value == "'=1+1"
    assert sheet["D1"].value == 18010
    csv_data, _, _ = export_books(BOOKS, LIST, "학교", format="csv")
    assert csv_data.startswith(b"\xef\xbb\xbf")
    assert "'=1+1" in csv_data.decode("utf-8-sig")
    html_data, _, _ = export_books([{**BOOKS[0], "title": "<script>x</script>"}], LIST, "학교", format="html")
    assert b"<script>x</script>" not in html_data
    assert b"&lt;script&gt;" in html_data


def test_template_footer_sum_expands_across_all_inserted_books(tmp_path: Path):
    path = tmp_path / "totals.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["도서명", "ISBN", "금액"])
    sheet["C3"] = "=SUM(C2:C2)"
    workbook.save(path)
    content, _, _ = export_books(BOOKS, LIST, "학교", template_path=path)
    exported = load_workbook(BytesIO(content)).active
    assert exported["C4"].value == "=SUM(C2:C3)"


def test_template_rejects_external_formula(tmp_path: Path):
    path = tmp_path / "external.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["도서명", "ISBN"])
    sheet["A3"] = '=WEBSERVICE("https://example.invalid/")'
    workbook.save(path)
    with pytest.raises(ValueError, match="외부"):
        inspect_template(path)
