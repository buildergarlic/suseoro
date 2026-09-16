from __future__ import annotations

from pathlib import Path
import zipfile

import pytest
from openpyxl import Workbook

from suseoro.simple.documents import parse_upload


def workbook_file(tmp_path: Path, rows, *, offset=0, column=0, merges=(), extra=None):
    path = tmp_path / "추천.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "추천"
    for row_number, row in enumerate(rows, offset + 1):
        for column_number, value in enumerate(row, column + 1):
            sheet.cell(row_number, column_number, value)
    for merged in merges:
        sheet.merge_cells(merged)
    if extra:
        sheet = workbook.create_sheet("추가")
        for row in extra:
            sheet.append(row)
    workbook.save(path)
    return path


def test_blank_leading_rows_and_columns_keep_real_source_and_formula(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가"], ["책", 9788937464010, "=10000+2000"]], offset=4, column=1)
    row = parse_upload(path, path.name)["rows"][0]
    assert row["provenance"]["row"] == 6
    assert row["raw_values"]["정가"] == "=10000+2000"
    assert row["price"] is None and row["needs_review"]


def test_merged_header_rows_form_single_book_with_price_and_quantity(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "신청정보", None], [None, None, "정가", "수량"], ["책", "9788937464010", 12000, 2]], merges=("A1:A2", "B1:B2", "C1:D1"))
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 1
    assert (result["rows"][0]["title"], result["rows"][0]["price"], result["rows"][0]["quantity"]) == ("책", 12000, 2)
    assert result["rows"][0]["provenance"]["row"] == 3


def test_late_header_repeated_header_and_totals_are_accounted_without_fake_books(tmp_path):
    headers = ["도서명", "ISBN", "정가", "수량"]
    path = workbook_file(tmp_path, [["ISBN"]] + [[f"안내 {i}"] for i in range(27)] + [headers, ["책", "9788937464010", 12000, 2], headers, ["다른 책", "9788936434267", 9000, 1], ["합계", None, 33000, 3]])
    result = parse_upload(path, path.name)
    assert [row["title"] for row in result["rows"]] == ["책", "다른 책"]
    assert [row["provenance"]["row"] for row in result["rows"]] == [30, 32]
    assert {item["kind"] for item in result["diagnostics"]} >= {"preamble", "repeated_header", "summary"}
    assert any(item["row"] == 33 and "합계" in item["raw_text"] for item in result["diagnostics"])


def test_mapping_one_sheet_does_not_erase_aliases_in_another_sheet(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가", "수량"], ["책", "9788937464010", 12000, 2]], extra=[["책제목", "도서번호", "가격", "권수"], ["다른 책", "9788936434267", 9000, 3]])
    result = parse_upload(path, path.name, mapping={"title": "도서명", "isbn": "ISBN", "price": "정가", "quantity": "수량"})
    assert [(row["title"], row["quantity"]) for row in result["rows"]] == [("책", 2), ("다른 책", 3)]


def test_mapping_can_explicitly_ignore_a_field(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가"], ["책", "9788937464010", 12000]])
    row = parse_upload(path, path.name, mapping={"price": ""})["rows"][0]
    assert row["title"] == "책" and row["price"] is None


def test_units_in_headers_and_quantity_are_read(tmp_path):
    path = workbook_file(tmp_path, [["도서명(서명)", "ISBN(13자리)", "정가(원)", "신청수량(권)"], ["책", "9788937464010", "12,000원", "2권"]])
    row = parse_upload(path, path.name)["rows"][0]
    assert (row["isbn"], row["price"], row["quantity"]) == ("9788937464010", 12000, 2)


@pytest.mark.parametrize("first", [None, "오류", "8937464012"])
def test_isbn13_is_used_when_isbn10_column_is_empty_invalid_or_equivalent(tmp_path, first):
    path = workbook_file(tmp_path, [["도서명", "ISBN10", "ISBN13", "정가"], ["책", first, "9788937464010", 12000]])
    result = parse_upload(path, path.name)
    row = result["rows"][0]
    assert row["isbn"] == "9788937464010"
    assert result["mapping"]["isbn"] == "ISBN13"
    assert row["raw_values"]["ISBN10"] in (first, "")


def test_conflicting_valid_isbn_columns_require_review(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "ISBN13", "정가"], ["책", "9788936434267", "9788937464010", 12000]])
    row = parse_upload(path, path.name)["rows"][0]
    assert row["needs_review"] and any("ISBN" in item and "서로" in item for item in row["warnings"])


@pytest.mark.parametrize("isbn", ["9.78893746401E+12", "9788937464010.0"])
def test_exact_numeric_isbn_text_is_normalized_without_changing_raw(tmp_path, isbn):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가"], ["책", isbn, 12000]])
    row = parse_upload(path, path.name)["rows"][0]
    assert row["isbn"] == "9788937464010"
    assert row["raw_values"]["ISBN"] == isbn


def test_rounded_or_bad_checksum_isbn_remains_for_review(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가"], ["책", "9.78894E+12", 12000], ["다른 책", "9788937464011", 10000]])
    rows = parse_upload(path, path.name)["rows"]
    assert [row["isbn"] for row in rows] == ["9.78894E+12", "9788937464011"]
    assert all(row["needs_review"] for row in rows)


@pytest.mark.parametrize("encoding", ["utf-16", "cp949"])
def test_korean_excel_unicode_or_legacy_text(tmp_path, encoding):
    path = tmp_path / "추천.txt"
    path.write_text("도서명\tISBN\t정가\t수량\n책\t9788937464010\t12000\t2\n", encoding=encoding)
    row = parse_upload(path, path.name)["rows"][0]
    assert (row["title"], row["isbn"], row["price"], row["quantity"]) == ("책", "9788937464010", 12000, 2)


def test_headerless_text_only_list_keeps_first_book(tmp_path):
    path = tmp_path / "목록.csv"
    path.write_text("첫 책,김작가,첫출판\n둘째 책,이작가,둘출판\n", encoding="utf-8")
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 2
    assert result["rows"][0]["raw_values"]["열 1"] == "첫 책"
    assert all(row["needs_review"] for row in result["rows"])


def test_headerless_first_book_without_isbn_or_price_is_not_consumed(tmp_path):
    path = workbook_file(tmp_path, [["가격없는책", "저자A", "출판사A", "", ""], ["가격있는책", "저자B", "출판사B", "9788937464010", 12000]])
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 2
    assert [row["raw_values"]["열 1"] for row in result["rows"]] == ["가격없는책", "가격있는책"]
    assert [row["provenance"]["row"] for row in result["rows"]] == [1, 2]
    assert all(row["needs_review"] for row in result["rows"])
    mapped = parse_upload(path, path.name, mapping={"title": "열 1", "author": "열 2", "publisher": "열 3", "isbn": "열 4", "price": "열 5"})
    assert [row["title"] for row in mapped["rows"]] == ["가격없는책", "가격있는책"]
    assert mapped["rows"][0]["price"] is None and mapped["rows"][0]["needs_review"]
    assert mapped["rows"][1]["price"] == 12000


def test_unknown_csv_mapping_keeps_quoted_newline(tmp_path):
    path = tmp_path / "목록.csv"
    path.write_text('A,B,C\n"두 줄\n제목",9788937464010,12000\n', encoding="utf-8", newline="")
    row = parse_upload(path, path.name, mapping={"title": "A", "isbn": "B", "price": "C"})["rows"][0]
    assert row["title"] == "두 줄\n제목"
    assert row["raw_values"]["A"] == "두 줄\n제목"


def test_recommendation_prose_with_price_comma_keeps_isbn(tmp_path):
    path = tmp_path / "추천.txt"
    path.write_text("책 제목 / 김작가 / 출판사 / 9788937464010 / 12,000원\n", encoding="utf-8")
    row = parse_upload(path, path.name)["rows"][0]
    assert row["isbn"] == "9788937464010"
    assert row["title"] == "" and row["needs_review"]
    assert "12,000원" in row["raw_text"]


def test_labeled_recommendation_blocks_become_reviewable_books(tmp_path):
    path = tmp_path / "추천.txt"
    path.write_text("도서명: 첫 책\n저자: 김작가\nISBN: 9788937464010\n정가: 12,000원\n\n도서명: 둘째 책\nISBN: 9788936434267\n정가: 9000원\n", encoding="utf-8")
    rows = parse_upload(path, path.name)["rows"]
    assert [(row["title"], row["isbn"], row["price"]) for row in rows] == [("첫 책", "9788937464010", 12000), ("둘째 책", "9788936434267", 9000)]
    assert rows[0]["provenance"]["row"] == 1 and rows[1]["provenance"]["row"] == 6
    assert all(row["needs_review"] and "ISBN:" in row["raw_text"] for row in rows)


def test_librarian_recommendation_template_preserves_all_bibliographic_columns(tmp_path):
    # Structure from the reported file; no dependency on the user's Drive.
    path = workbook_file(tmp_path, [[], [None, "어린이도서연구회 / 독서단체 추천도서"], [], ["번호", "도서 제목", "지은이", "출판사", "ISBN", "정가(원)", "판형", "가로(mm)", "세로(mm)", "쪽수", "출간일", "분야(KDC)", "추천기관", "수상내역"], [1, "첫 책", "김작가 ", "출판사 하나", 9788943318352, 16000, None, None, None, 48, "2025-11-14", "문학", None, None], [2, "둘째 책", "이작가", "출판사 둘", 9791165736811, 12000, None, None, None, 32, "2025-09-26", "문학", None, None]])
    result = parse_upload(path, path.name)
    assert len(result["rows"]) == 2
    first, second = result["rows"]
    assert (first["title"], first["author"], first["publisher"], first["isbn"], first["price"], first["category"], first["published_date"]) == ("첫 책", "김작가", "출판사 하나", "9788943318352", 16000, "문학", "2025-11-14")
    assert first["raw_values"]["지은이"] == "김작가 "
    assert [row["provenance"]["row"] for row in result["rows"]] == [5, 6]
    assert sum(row["price"] for row in result["rows"]) == 28000
    assert all(not row["needs_review"] for row in result["rows"])


def test_multiple_prose_recommendations_keep_isbn_and_full_price(tmp_path):
    path = tmp_path / "추천.txt"
    path.write_text("첫 책 / 김작가 / 9788937464010 / 12,000원\n둘째 책 / 이작가 / 9788936434267 / 9,000원\n", encoding="utf-8")
    rows = parse_upload(path, path.name)["rows"]
    assert [row["isbn"] for row in rows] == ["9788937464010", "9788936434267"]
    assert "12,000원" in rows[0]["raw_text"]


def test_docx_labeled_paragraphs_are_grouped_into_books(tmp_path):
    path = tmp_path / "추천.docx"
    paragraphs = ["도서명: 첫 책", "저자: 김작가", "ISBN: 9788937464010", "정가: 12000", "", "도서명: 둘째 책", "ISBN: 9788936434267", "정가: 9000"]
    xml = '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>' + "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in paragraphs) + '</w:body></w:document>'
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", xml)
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    rows = parse_upload(path, path.name)["rows"]
    assert [(row["title"], row["isbn"], row["price"]) for row in rows] == [("첫 책", "9788937464010", 12000), ("둘째 책", "9788936434267", 9000)]
    assert [row["provenance"]["row"] for row in rows] == [1, 6]


def test_explicit_unmapped_header_does_not_hide_a_quantity_problem(tmp_path):
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가", "신청수량(권)"], ["책", "9788937464010", 12000, "여러 권"]])
    row = parse_upload(path, path.name)["rows"][0]
    assert row["quantity"] == 1 and row["needs_review"]
    assert any("수량" in warning for warning in row["warnings"])


@pytest.mark.parametrize("header", ["link", "참고 링크", "도서정보URL"])
def test_reference_link_column_is_preserved(tmp_path, header):
    url = "https://library.example/books/9788937464010?source=추천"
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가", header], ["책", "9788937464010", 12000, url]])
    result = parse_upload(path, path.name)
    assert result["rows"][0]["link"] == url
    assert result["rows"][0]["raw_values"][header] == url
    assert result["mapping"]["link"] == header


def test_reference_link_can_be_mapped_from_custom_column_and_ignored(tmp_path):
    url = "https://library.example/books/9788937464010"
    path = workbook_file(tmp_path, [["도서명", "ISBN", "정가", "외부자료"], ["책", "9788937464010", 12000, url]])
    mapped = parse_upload(path, path.name, mapping={"link": "외부자료"})
    assert mapped["rows"][0]["link"] == url
    ignored = parse_upload(path, path.name, mapping={"link": ""})
    assert ignored["rows"][0]["link"] == ""
    assert ignored["rows"][0]["raw_values"]["외부자료"] == url
