from __future__ import annotations

import hashlib
import struct
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from python_calamine import CalamineWorkbook

from suseoro.ingestion.contracts import DocumentRole, RowStatus
from suseoro.ingestion.parsers.tabular import _parse_sheets, parse_tabular
from suseoro.ingestion.parsers.text import parse_pasted_text

UPSTREAM_XLSB_SHA256 = (
    "b3a6b5034076373fcf2c93979839738275fac80b9de64ac557235f5a286f7e89"
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _parse_file(path: Path, role: DocumentRole):
    return parse_tabular(path, role=role, sha256=_sha256(path))


def _make_xlsx(path: Path, *, rows: int = 1) -> None:
    from openpyxl import Workbook

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("도서")
    sheet.append(["ISBN", "제목", "수량"])
    for number in range(rows):
        isbn = "0012345678901" if number == 0 else f"978000000{number:04d}"
        sheet.append([isbn, f"책 {number}", number + 1])
    workbook.save(path)


def _make_complex_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "견적"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "2026년 추천 도서 견적"
    sheet.append([])
    sheet.append(["ISBN", "제목", "제목", "수량", "단가"])
    sheet.append(["0012345678901", "원제", "부제", 2, "=1000+200"])
    sheet.append([])
    sheet.append([None, "행 오류", None, 1, 500])
    other = workbook.create_sheet("추가")
    other.append(["ISBN", "제목", "수량", "단가"])
    other.append(["0099999999999", "추가 책", 1, 3000])
    workbook.save(path)


def _make_xls(path: Path) -> None:
    import xlwt

    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Books")
    for column, value in enumerate(["ISBN", "제목", "수량"]):
        sheet.write(0, column, value)
    for column, value in enumerate(["0012345678901", "옛 책", 2]):
        sheet.write(1, column, value)
    workbook.save(str(path))


def _make_ods(path: Path) -> None:
    from odf.opendocument import OpenDocumentSpreadsheet
    from odf.table import Table, TableCell, TableRow
    from odf.text import P

    document = OpenDocumentSpreadsheet()
    table = Table(name="Books")
    for values in (["ISBN", "제목", "수량"], ["0012345678901", "열린 책", "2"]):
        row = TableRow()
        for value in values:
            cell = TableCell(valuetype="string")
            cell.addElement(P(text=value))
            row.addElement(cell)
        table.addElement(row)
    document.spreadsheet.addElement(table)
    document.save(str(path))


def _make_minimal_xlsb_package(path: Path) -> None:
    """Construct an XLSB-signature OPC package; parser must reject it safely, not mislabel it."""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            "application/vnd.ms-excel.sheet.binary.macroEnabled.main",
        )
        archive.writestr("xl/workbook.bin", b"\x00\x01not-a-real-biff12-stream")


def _make_data_bearing_xlsb(path: Path) -> None:
    """Build a deterministic BIFF12 workbook based on calamine's public reproducer."""

    def record(record_id: int, payload: bytes = b"") -> bytes:
        output = bytearray()
        if record_id < 0x80:
            output.append(record_id)
        else:
            output += bytes([(record_id & 0x7F) | 0x80, (record_id >> 7) & 0xFF])
        size = len(payload)
        while True:
            byte = size & 0x7F
            size >>= 7
            output.append(byte | 0x80 if size else byte)
            if not size:
                break
        return bytes(output) + payload

    def wide_string(value: str) -> bytes:
        return struct.pack("<I", len(value)) + value.encode("utf-16-le")

    workbook = record(0x0083) + record(0x0099, b"\x00" * 4)
    workbook += record(0x008F)
    workbook += record(
        0x009C,
        struct.pack("<II", 0, 0) + wide_string("rId1") + wide_string("Books"),
    )
    workbook += record(0x0090) + record(0x0084)

    def string_cell(column: int, value: str) -> bytes:
        return record(0x0006, struct.pack("<II", column, 0) + wide_string(value))

    def number_cell(column: int, value: float) -> bytes:
        return record(0x0005, struct.pack("<II", column, 0) + struct.pack("<d", value))

    sheet = (
        record(0x0081)
        + record(0x0094, struct.pack("<IIII", 0, 1, 0, 3))
        + record(0x0091)
        + record(0x0000, struct.pack("<I", 0))
        + string_cell(0, "ISBN")
        + string_cell(1, "제목")
        + string_cell(2, "수량")
        + string_cell(3, "단가")
        + record(0x0000, struct.pack("<I", 1))
        + string_cell(0, "0012345678901")
        + string_cell(1, "바이너리 책")
        + number_cell(2, 2.0)
        + number_cell(3, 1234.5)
        + record(0x0092)
        + record(0x0082)
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="bin" ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/><Override PartName="/xl/worksheets/sheet1.bin" ContentType="application/vnd.ms-excel.worksheet"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.bin"/></Relationships>',
        )
        archive.writestr("xl/workbook.bin", workbook)
        archive.writestr(
            "xl/_rels/workbook.bin.rels",
            '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.bin"/></Relationships>',
        )
        archive.writestr("xl/worksheets/sheet1.bin", sheet)


@pytest.mark.parametrize(
    ("suffix", "builder", "expected_format"),
    [
        (".xlsx", _make_xlsx, "XLSX"),
        (".xls", _make_xls, "XLS"),
        (".ods", _make_ods, "ODS"),
    ],
)
def test_calamine_parses_real_spreadsheet_formats(
    tmp_path: Path, suffix: str, builder: object, expected_format: str
) -> None:
    """Bypassing calamine or losing identifier text in a supported workbook must fail."""
    target = tmp_path / f"books{suffix}"
    builder(target)  # type: ignore[operator]

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert result.detected_format == expected_format
    assert result.parser_backend == "calamine"
    assert [row.status for row in result.rows] == [RowStatus.SUCCESS]
    assert result.rows[0].fields["isbn"].value == "0012345678901"
    assert result.rows[0].provenance.sheet == ("도서" if suffix == ".xlsx" else "Books")


def test_xlsb_signature_is_recognized_and_broken_payload_is_accounted_for(
    tmp_path: Path,
) -> None:
    """Mislabeling XLSB or silently returning no outcome for a broken sheet must fail."""
    target = tmp_path / "broken.xlsb"
    _make_minimal_xlsb_package(target)

    result = _parse_file(target, DocumentRole.UNKNOWN)

    assert result.detected_format == "XLSB"
    assert result.rows[0].status == RowStatus.ROW_ERROR
    assert result.rows[0].error_code == "WORKBOOK_ERROR"


def test_calamine_parses_data_bearing_xlsb_with_complete_provenance(
    tmp_path: Path,
) -> None:
    """An empty-only XLSB path or missing values/provenance must fail this test."""
    target = tmp_path / "data-bearing.xlsb"
    _make_data_bearing_xlsb(target)
    digest = _sha256(target)

    result = _parse_file(target, DocumentRole.UNKNOWN)

    assert result.detected_format == "XLSB"
    assert result.parser_backend == "calamine"
    assert result.header_rows == {"Books": 1}
    assert [row.status for row in result.rows] == [RowStatus.SUCCESS]
    row = result.rows[0]
    assert row.fields["isbn"].value == "0012345678901"
    assert row.fields["title"].value == "바이너리 책"
    assert row.fields["quantity"].value == 2.0
    assert row.fields["unit_price"].value == 1234.5
    assert (
        row.provenance.source_file_sha256,
        row.provenance.sheet,
        row.provenance.source_row,
    ) == (
        digest,
        "Books",
        2,
    )


def test_redistributable_upstream_xlsb_fixture_contains_real_typed_values() -> None:
    """Replacing the licensed XLSB corpus with an empty/signature-only file must fail."""
    fixture = (
        Path(__file__).parent / "fixtures" / "ingestion" / "python-calamine-base.xlsb"
    )
    assert _sha256(fixture) == UPSTREAM_XLSB_SHA256

    workbook = CalamineWorkbook.from_path(fixture)
    try:
        values = workbook.get_sheet_by_name("Sheet1").to_python()[0]
    finally:
        workbook.close()

    assert values[:5] == ["String", 1.0, 1.1, True, False]


def test_csv_tsv_txt_and_paste_detect_encoding_delimiter_and_stream_rows(
    tmp_path: Path,
) -> None:
    """Assuming UTF-8/comma or using a separate paste contract must fail this test."""
    csv_path = tmp_path / "bom.data"
    csv_path.write_bytes("\ufeffISBN,제목,수량\n00123,쉼표 책,2\n".encode("utf-8"))
    tsv_path = tmp_path / "korean.data"
    tsv_path.write_bytes("ISBN\t제목\t수량\n00456\t탭 책\t3\n".encode("cp949"))
    txt_path = tmp_path / "pipe.data"
    txt_path.write_text("ISBN|제목|수량\n00789|파이프 책|4\n", encoding="utf-8")

    parsed = [
        _parse_file(csv_path, DocumentRole.PURCHASE_REQUEST),
        _parse_file(tsv_path, DocumentRole.PURCHASE_REQUEST),
        _parse_file(txt_path, DocumentRole.PURCHASE_REQUEST),
        parse_pasted_text(
            "ISBN\t제목\t수량\n00999\t붙여넣기 책\t5\n",
            role=DocumentRole.PURCHASE_REQUEST,
        ),
    ]

    assert [item.detected_format for item in parsed] == ["CSV", "TSV", "TXT", "TSV"]
    assert [item.rows[0].fields["isbn"].value for item in parsed] == [
        "00123",
        "00456",
        "00789",
        "00999",
    ]
    assert parsed[1].encoding in {"cp949", "euc-kr"}
    assert parsed[3].source_kind == "PASTE"
    assert all(
        row.provenance.source_file_sha256 for item in parsed for row in item.rows
    )


@pytest.mark.parametrize(
    ("encoding", "character"),
    [("utf-8", "한"), ("cp949", "한")],
)
def test_text_probe_accepts_multibyte_character_split_at_64k_boundary(
    tmp_path: Path, encoding: str, character: str
) -> None:
    """Finalizing the 64 KiB probe decoder mid-character must fail this test."""
    target = tmp_path / f"boundary-{encoding}.csv"
    prefix = "ISBN,제목\n00123,".encode(encoding)
    filler = b"a" * (65_535 - len(prefix))
    target.write_bytes(prefix + filler + character.encode(encoding) + b"\n")

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert result.rows[0].status == RowStatus.SUCCESS
    assert result.rows[0].fields["title"].value.endswith(character)
    assert result.encoding in (
        {"utf-8", "utf-8-sig"} if encoding == "utf-8" else {"cp949", "euc-kr"}
    )


def test_late_invalid_text_bytes_become_accounted_row_error(tmp_path: Path) -> None:
    """Aborting and losing earlier rows on a late decode error must fail this test."""
    target = tmp_path / "late-invalid.csv"
    first_title = b"a" * 70_000
    target.write_bytes(b"ISBN,title\n00123," + first_title + b"\n00456,broken-\xff\n")

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert [row.status for row in result.rows] == [
        RowStatus.SUCCESS,
        RowStatus.ROW_ERROR,
    ]
    assert result.rows[1].error_code == "DECODE_ERROR"
    assert [row.provenance.source_row for row in result.rows] == [2, 3]


def test_text_row_iterator_is_normalized_before_the_next_row_is_read() -> None:
    """Materializing the whole text source before row normalization must fail."""
    events: list[str] = []

    class Identifier:
        def __str__(self) -> str:
            assert "row-after-probe-read" not in events
            return "00123"

    def source_rows() -> object:
        yield ["ISBN", "제목"]
        yield [Identifier(), "첫 책"]
        for number in range(23):
            yield [f"978{number:010d}", f"책 {number}"]
        events.append("row-after-probe-read")
        yield ["00456", "둘째 책"]

    result = _parse_sheets(
        [("텍스트", source_rows())],  # type: ignore[arg-type]
        role=DocumentRole.PURCHASE_REQUEST,
        detected_format="CSV",
        parser_backend="python-csv-stream",
        source_file_sha256="a" * 64,
    )

    assert result.rows[0].fields["isbn"].value == "00123"
    assert result.rows[-1].fields["isbn"].value == "00456"


def test_complex_xlsx_handles_notes_merge_sheets_blanks_duplicates_and_formulas(
    tmp_path: Path,
) -> None:
    """Dropping a logical row/sheet or exposing formula source as evaluated data must fail."""
    target = tmp_path / "complex.xlsx"
    _make_complex_xlsx(target)

    result = _parse_file(target, DocumentRole.VENDOR_QUOTE)

    assert result.header_rows == {"견적": 3, "추가": 1}
    assert [row.provenance.sheet for row in result.rows] == ["견적", "견적", "추가"]
    assert [row.provenance.source_row for row in result.rows] == [4, 6, 2]
    assert [row.status for row in result.rows] == [
        RowStatus.SUCCESS,
        RowStatus.ROW_ERROR,
        RowStatus.SUCCESS,
    ]
    assert list(result.rows[0].raw_values) == [
        "ISBN",
        "제목",
        "제목__2",
        "수량",
        "단가",
    ]
    assert result.rows[0].fields["isbn"].value == "0012345678901"
    assert result.rows[0].fields["unit_price"].value is None
    assert result.rows[0].fields["unit_price"].raw_value == "=1000+200"
    assert "FORMULA_LIKE_INPUT" in {
        warning.code for warning in result.rows[0].fields["unit_price"].warnings
    }
    logical_keys = {
        (
            row.provenance.source_file_sha256,
            row.provenance.sheet,
            row.provenance.source_row,
        )
        for row in result.rows
    }
    assert len(logical_keys) == len(result.rows)


@pytest.mark.parametrize("marker", ["=", "+", "-", "@"])
def test_formula_like_text_is_preserved_and_warned_without_execution(
    tmp_path: Path, marker: str
) -> None:
    """Losing or silently accepting a formula-like mapped/unmapped string must fail."""
    target = tmp_path / f"formula-like-{ord(marker)}.csv"
    title = marker + "title-expression"
    note = marker + "note-expression"
    target.write_text(f"ISBN,제목,메모\n00123,{title},{note}\n", encoding="utf-8")

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    row = result.rows[0]
    assert row.fields["title"].value == title
    assert row.fields["title"].raw_value == title
    assert "FORMULA_LIKE_INPUT" in {
        warning.code for warning in row.fields["title"].warnings
    }
    assert "FORMULA_LIKE_INPUT" in {warning.code for warning in row.warnings}


def _swap_xlsx_worksheet_relationships(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        members = {
            info.filename: source.read(info.filename) for info in source.infolist()
        }
    relationships_name = "xl/_rels/workbook.xml.rels"
    root = ElementTree.fromstring(members[relationships_name])
    worksheet_relationships = [
        element
        for element in root
        if element.attrib.get("Type", "").endswith("/worksheet")
    ]
    first_target = worksheet_relationships[0].attrib["Target"]
    worksheet_relationships[0].attrib["Target"] = worksheet_relationships[1].attrib[
        "Target"
    ]
    worksheet_relationships[1].attrib["Target"] = first_target
    members[relationships_name] = ElementTree.tostring(
        root, encoding="utf-8", xml_declaration=True
    )
    replacement = path.with_suffix(".replacement.xlsx")
    with zipfile.ZipFile(replacement, "w", compression=zipfile.ZIP_DEFLATED) as target:
        for name, contents in members.items():
            target.writestr(name, contents)
    replacement.replace(path)


def test_xlsx_formula_relationships_and_unmapped_columns_preserve_provenance(
    tmp_path: Path,
) -> None:
    """Assuming sheetN order or scanning mapped columns only must fail this test."""
    from openpyxl import Workbook

    target = tmp_path / "relationship-order.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["ISBN", "제목"])
    first.append(["00111", "plain"])
    second = workbook.create_sheet("Second")
    second.append(["ISBN", "제목", "메모"])
    second.append(["00222", "formula source", "=1+2"])
    workbook.save(target)
    _swap_xlsx_worksheet_relationships(target)

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    by_sheet = {row.provenance.sheet: row for row in result.rows}
    formula_row = by_sheet["First"]
    assert formula_row.raw_values["메모"] == "=1+2"
    assert "FORMULA_LIKE_INPUT" in {warning.code for warning in formula_row.warnings}
    assert not by_sheet["Second"].warnings


def test_broken_workbook_returns_explicit_workbook_error_row(tmp_path: Path) -> None:
    """Raising an opaque exception or returning zero accounted rows must fail this test."""
    target = tmp_path / "broken.xlsx"
    target.write_bytes(b"PK\x03\x04broken")

    result = _parse_file(target, DocumentRole.UNKNOWN)

    assert len(result.rows) == 1
    assert result.rows[0].status == RowStatus.ROW_ERROR
    assert result.rows[0].error_code == "WORKBOOK_ERROR"


def test_parse_requires_canonical_sha_and_populates_every_error_provenance(
    tmp_path: Path,
) -> None:
    """Optional/malformed SHA provenance or a synthetic error without SHA must fail."""
    target = tmp_path / "broken.xlsx"
    target.write_bytes(b"PK\x03\x04broken")

    with pytest.raises(ValueError, match="sha256"):
        parse_tabular(target, role=DocumentRole.UNKNOWN, sha256="NOT-A-HASH")
    with pytest.raises(ValueError, match="sha256"):
        parse_tabular(target, role=DocumentRole.UNKNOWN)

    digest = _sha256(target)
    result = parse_tabular(target, role=DocumentRole.UNKNOWN, sha256=digest)
    assert result.rows[0].provenance.source_file_sha256 == digest


def test_compatible_xlsx_uses_read_only_data_only_fallback_without_formula_execution(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Falling back for non-XLSX or loading formulas/entire workbook must fail this test."""
    target = tmp_path / "compat.xlsx"
    _make_xlsx(target)

    from suseoro.ingestion.parsers import tabular

    real_loader = tabular._load_with_openpyxl
    calls: list[tuple[bool, bool]] = []

    def rejecting_calamine(path: Path) -> object:
        raise tabular.CalamineCompatibilityError("compatibility edge")

    def observing_loader(path: Path, *, read_only: bool, data_only: bool) -> object:
        calls.append((read_only, data_only))
        return real_loader(path, read_only=read_only, data_only=data_only)

    monkeypatch.setattr(tabular, "_load_with_calamine", rejecting_calamine)
    monkeypatch.setattr(tabular, "_load_with_openpyxl", observing_loader)

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert result.parser_backend == "openpyxl-read-only-data-only"
    assert calls == [(True, True)]
    assert result.rows[0].status == RowStatus.SUCCESS


def test_data_cells_beyond_header_width_remain_in_raw_provenance(
    tmp_path: Path,
) -> None:
    """Truncating unexpected source cells from raw row provenance must fail this test."""
    from openpyxl import Workbook

    target = tmp_path / "extra-cell.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["ISBN", "제목"])
    sheet.append(["00123", "책", "unexpected"])
    workbook.save(target)

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert result.rows[0].status == RowStatus.SUCCESS
    assert result.rows[0].raw_values == {
        "ISBN": "00123",
        "제목": "책",
        "열3": "unexpected",
    }


def test_twenty_thousand_row_fixture_is_complete_and_reports_timing(
    tmp_path: Path,
) -> None:
    """A parser that truncates the large fixture or omits timing must fail this test."""
    target = tmp_path / "large.xlsx"
    _make_xlsx(target, rows=20_000)

    result = _parse_file(target, DocumentRole.PURCHASE_REQUEST)

    assert len(result.rows) == 20_000
    assert result.elapsed_seconds > 0
